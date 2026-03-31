# 記得先在終端機安裝：pip install python-pptx
from pptx import Presentation
from pptx.util import Inches

# 建立簡報物件
prs = Presentation()
blank_slide_layout = prs.slide_layouts[5] # 只有標題的版面
slide = prs.slides.add_slide(blank_slide_layout)

# 設定標題
title_shape = slide.shapes.title
title_shape.text = "Limitations of Proactive Defenses (Adversarial Training) in NTP"

# 痛點 1：Trade-off
txBox = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(3), Inches(3))
tf = txBox.text_frame
tf.text = "1. Accuracy–Robustness Trade-off"
p = tf.add_paragraph()
p.text = "• High robustness on attacked data (e.g., 12%)"
p = tf.add_paragraph()
p.text = "• Low accuracy on clean data (Clean traffic is the majority)"

# 痛點 2：Weak Generalization
txBox2 = slide.shapes.add_textbox(Inches(3.5), Inches(2), Inches(3), Inches(3))
tf2 = txBox2.text_frame
tf2.text = "2. Weak Generalization"
p2 = tf2.add_paragraph()
p2.text = "• Train 12% -> Test 3% (Mismatched): Accuracy Drop"
p2 = tf2.add_paragraph()
p2.text = "• Train 3% -> Test 12% (Defense Failure): Vulnerable to Attack"
p2 = tf2.add_paragraph()
p2.text = "• Fails against real-world unknown perturbations"

# 痛點 3：No Online Adaptation
txBox3 = slide.shapes.add_textbox(Inches(6.5), Inches(2), Inches(3), Inches(3))
tf3 = txBox3.text_frame
tf3.text = "3. No Online Adaptation"
p3 = tf3.add_paragraph()
p3.text = "• Defenses remain static after deployment"
p3 = tf3.add_paragraph()
p3.text = "• Offline Static Training -> Dynamic Online Phase"
p3 = tf3.add_paragraph()
p3.text = "• Inability to Auto-adapt"

# 底部 Motivation
txBox4 = slide.shapes.add_textbox(Inches(0.5), Inches(6), Inches(9), Inches(1))
tf4 = txBox4.text_frame
tf4.text = "Motivation: Especially in NTP, these three critical gaps highlight the need for our proposed approach."

# 儲存檔案
prs.save('Limitations_of_Proactive_Defenses.pptx')
print("PPTX 檔案生成完成！")